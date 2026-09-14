import tkinter as tk

from pathlib import Path
from tkinter import filedialog

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE = 5 * 1024 * 1024

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


def select_multiple_resumes(
    root: tk.Tk,
) -> tuple[str, ...]:

    return filedialog.askopenfilenames(
        parent=root,
        title="Select Resume Files",
        filetypes=[
            ("Supported files", "*.pdf *.docx"),
            ("PDF files", "*.pdf"),
            ("Word documents", "*.docx"),
            ("All files", "*.*"),
        ],
    )


def select_jd(
    root: tk.Tk,
) -> str:

    return filedialog.askopenfilename(
        parent=root,
        title="Select Job Description",
        filetypes=[
            ("Supported files", "*.pdf *.docx"),
            ("PDF files", "*.pdf"),
            ("Word documents", "*.docx"),
            ("All files", "*.*"),
        ],
    )


def select_input_files():

    root = tk.Tk()
    root.withdraw()

    try:

        print("\nSelect your resume file(s)...")

        resume_paths = select_multiple_resumes(
            root
        )

        if not resume_paths:
            raise ValueError(
                "No resume files were selected."
            )

        print(
            f"\n{len(resume_paths)} resume(s) selected."
        )

        print("\nSelect your job description...")

        jd_path = select_jd(root)

        if not jd_path:
            raise ValueError(
                "Job description selection was cancelled."
            )

        return resume_paths, jd_path

    finally:
        root.destroy()


def validate_file(
    file_path: str,
    file_label: str,
) -> Path:

    path = Path(file_path).expanduser()

    if not path.exists():
        raise FileNotFoundError(
            f"{file_label} does not exist: {path}"
        )

    if not path.is_file():
        raise ValueError(
            f"{file_label} is not a file: {path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported {file_label} format: {extension}"
        )

    size = path.stat().st_size

    if size > MAX_FILE_SIZE:
        raise ValueError(
            f"{file_label} exceeds the 5 MB limit."
        )

    if size == 0:
        raise ValueError(
            f"{file_label} is empty."
        )

    return path


def extract_pdf_text(
    path: Path,
) -> str:

    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:

        text = page.extract_text() or ""

        if text.strip():
            pages.append(text)

    result = "\n\n".join(pages).strip()

    if not result:
        raise ValueError(
            f"No readable text found in PDF: {path}"
        )

    return result


def extract_docx_text(
    path: Path,
) -> str:

    document = Document(str(path))

    sections = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                sections.append(
                    " | ".join(cells)
                )

    result = "\n".join(sections).strip()

    if not result:
        raise ValueError(
            f"No readable text found in DOCX: {path}"
        )

    return result


def extract_text(
    path: Path,
) -> str:

    if path.suffix.lower() == ".pdf":
        return extract_pdf_text(path)

    if path.suffix.lower() == ".docx":
        return extract_docx_text(path)

    raise ValueError(
        f"Unsupported file extension: "
        f"{path.suffix}"
    )